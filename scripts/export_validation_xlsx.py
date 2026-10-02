"""Validation workbook: latest-quarter metrics for every active lender, same-quarter history, data dictionary, coverage."""
import os, numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

D = os.path.join(os.path.dirname(__file__), '..', 'data'); OUT = '/mnt/user-data/outputs/lender_metrics_26Q2.xlsx'; OUT2 = '/mnt/user-data/outputs/lender_metrics_history_2016_2026.xlsx'
m = pd.read_parquet(f'{D}/metrics_quarterly.parquet'); L = pd.read_parquet(f'{D}/lenders.parquet'); B = pd.read_parquet(f'{D}/branches.parquet')
latest = m['period'].max(); lq = latest.quarter
M = 1e6

# (output column, source column, kind)  kind: txt, musd, pct(stored as fraction), num, flag
COLS = [
    ('Lender ID', 'lender_id', 'txt'), ('Name', 'name', 'txt'), ('Charter', 'charter_lbl', 'txt'), ('Charter class', 'charter_class', 'txt'), ('State', 'state', 'txt'), ('City', 'city', 'txt'),
    ('Community bank (FDIC)', 'community_bank', 'flag'), ('Field of membership (CU)', 'fom_class', 'txt'), ('Low-income designated (CU)', 'licu_flag', 'flag'), ('Period', 'period_lbl', 'txt'),
    ('Total Deposits ($M)', 'deposits', 'musd'), ('Total Assets ($M)', 'assets', 'musd'), ('Total Loans ($M)', 'loans', 'musd'),
    ('Loan to Deposit (%)', 'loan_to_deposit_pct', 'pct'), ('Tier 1 Leverage (%)', 'tier1_leverage_pct', 'pct'), ('Net Worth Ratio (%)', 'net_worth_ratio_pct', 'pct'),
    ('Funding Cost (%) [int exp / avg deposits+borrowings]', 'funding_cost_pct', 'pct'), ('FDIC-published cost of funding earning assets (%)', 'cost_funds_pct_published', 'pct'),
    ('NIM (%) [NII / avg assets]', 'nim_pct', 'pct'), ('FDIC-published NIM (%)', 'nim_pct_published', 'pct'),
    ('Investment CRE as % of Capital', 'cre_pct_capital', 'pct'), ('Construction as % of Capital', 'constr_pct_capital', 'pct'),
    ('Non-Owner-Occupied Nonres CRE ($M)', 'cre_nonoo', 'musd'), ('Multifamily ($M)', 'cre_multi', 'musd'), ('Construction & Land ($M)', 'cre_constr', 'musd'),
    ('Total Investment CRE ($M) [nonOO+multi+constr]', 'cre_investment', 'musd'), ('Owner-Occupied Nonres CRE ($M)', 'cre_oo', 'musd'), ('Total CRE ($M)', 'cre_total', 'musd'),
    ('Brokered Deposits % (banks)', 'brokered_pct', 'pct'), ('Non-member Deposits % (CUs)', 'nonmember_dep_pct', 'pct'), ('Uses brokered deposits', 'brokered_flag', 'flag'),
    ('CRE Net Change YTD ($M) [total CRE]', 'cre_net_change_ytd', 'musd'), ('Investment CRE Net Change YTD ($M)', 'cre_investment_net_change_ytd', 'musd'),
    ('CRE Granted YTD ($M) [CU only]', 'cre_granted_ytd', 'musd'), ('CRE Granted YTD (# loans) [CU only]', 'cre_granted_ytd_n', 'num'), ('CRE Granted YTD Avg Loan ($M) [CU only]', 'cre_granted_avg_loan', 'musd'),
    ('Remaining CRE Cap ($M) [bank: 300% of capital]', 'cre_cap_remaining', 'musd'), ('Remaining Construction Cap ($M) [bank: 100% of capital]', 'constr_cap_remaining', 'musd'),
    ('MBL Cap Used (%) [CU, non-LICU]', 'mbl_cap_used_pct', 'pct'), ('Remaining MBL Cap ($M) [CU, non-LICU]', 'mbl_cap_remaining', 'musd'),
    ('Reported CRE Interest Rate (%) [CU only]', 'cre_rate_pct', 'pct'), ('Blended Loan Yield (%)', 'loan_yield_pct', 'pct'),
    ('Non-Performing Loans % [bank 90+/nonaccrual; CU 60+]', 'npl_pct', 'pct'), ('CRE Non-Performing % [bank only]', 'cre_npl_pct', 'pct'), ('ROA (%)', 'roa_pct', 'pct'),
    ('Est. Legal Lending Limit ($M) [15% of capital]', 'lending_limit_est', 'musd'), ('Capital Base ($M)', 'total_rbc', 'musd'), ('Capital base type', 'capital_type', 'txt'),
]

def prep(df):
    df = df.merge(L[['lender_id', 'city', 'charter_class', 'fom_class', 'community_bank', 'active']].rename(columns={'community_bank': 'cb_dir'}), on='lender_id', how='left')
    df['charter_lbl'] = df['charter'].map({'bank': 'Bank', 'cu': 'Credit Union'})
    df['period_lbl'] = df['period'].dt.strftime('%y') + 'Q' + df['quarter'].astype(str)
    df['licu_flag'] = np.where(df['charter'] == 'cu', df['licu'] == 1, np.nan)
    df['community_bank'] = np.where(df['charter'] == 'bank', df['community_bank'] == 1, np.nan)
    df['capital_type'] = np.select([df['charter'] == 'cu', df['cblr_bank'] == 1], ['Net worth', 'Tier 1 + allowance (CBLR filer)'], 'Total risk-based capital')
    out = pd.DataFrame()
    for name, src, kind in COLS:
        v = df[src]
        if kind == 'musd': v = v / M
        elif kind == 'pct': v = v / 100.0
        elif kind == 'flag': v = v.map({True: 'Y', False: 'N', 1: 'Y', 0: 'N', 1.0: 'Y', 0.0: 'N'})
        out[name] = v
    return out

def style(ws, kinds, freeze='C2', widths=None):
    hdr_fill = PatternFill('solid', start_color='1F3864'); hdr_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    for c in ws[1]:
        c.fill = hdr_fill; c.font = hdr_font; c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 48
    fmt = {'musd': '#,##0.0;(#,##0.0);-', 'pct': '0.00%;(0.00%);-', 'num': '#,##0;(#,##0);-'}
    for j, k in enumerate(kinds, start=1):
        col = get_column_letter(j)
        ws.column_dimensions[col].width = (widths or {}).get(j, 14 if k != 'txt' else 18)
        if k in fmt:
            for cell in ws[col][1:]: cell.number_format = fmt[k]; cell.font = Font(name='Arial', size=10)
        else:
            for cell in ws[col][1:]: cell.font = Font(name='Arial', size=10)
    ws.freeze_panes = freeze; ws.auto_filter.ref = ws.dimensions

# ---- sheet 1: latest quarter, active lenders
act = set(L[L['active']]['lender_id'])
lat = m[(m['period'] == latest) & (m['lender_id'].isin(act))].copy()
lat = lat.merge(L[['lender_id', 'charter_class']], on='lender_id', how='left', suffixes=('', '_y'))
lat = lat[lat['charter_class'] != 'Insured branch of foreign bank'].drop(columns=['charter_class'])
S1 = prep(lat).sort_values(['Charter', 'Total Assets ($M)'], ascending=[True, False])

# ---- sheet 2: same-quarter history 2016..latest for all lenders that are active today
hist = m[(m['quarter'] == lq) & (m['lender_id'].isin(act))].copy()
S2 = prep(hist).sort_values(['Charter', 'Name', 'Period'])

# ---- sheet 3: coverage
cov = m.groupby(['period', 'charter']).agg(n=('lender_id', 'size'), assets_bn=('assets', lambda s: s.sum() / 1e9), cre_total_bn=('cre_total', lambda s: s.sum() / 1e9),
                                           cre_investment_bn=('cre_investment', lambda s: s.sum() / 1e9), cre_granted_bn=('cre_granted_ytd', lambda s: s.sum() / 1e9),
                                           cre_rate_reported=('cre_rate_pct', lambda s: s.notna().sum())).reset_index()
cov['period'] = cov['period'].dt.date
cov.columns = ['Period', 'Charter', 'Institutions', 'Total assets ($B)', 'Total CRE ($B)', 'Investment CRE ($B)', 'CRE granted YTD ($B, CU)', 'CUs reporting a CRE rate']

# ---- sheet 4: branches for one state (Illinois) as a structure sample
S4 = B[B['state'] == 'IL'].merge(L[['lender_id', 'name', 'charter']].rename(columns={'name': 'lender_name'}), on='lender_id', how='left')
S4 = S4[['branch_id', 'lender_id', 'lender_name', 'charter', 'name', 'site_type', 'main_office', 'address', 'city', 'state', 'zip', 'county', 'lat', 'lon'] + [c for c in B.columns if c.startswith('dep_')]].copy()
for c in [c for c in S4.columns if c.startswith('dep_')]: S4[c] = S4[c] / M
S4 = S4.rename(columns={c: f"Deposits {c[4:]} ($M, SOD June 30)" for c in S4.columns if c.startswith('dep_')})
S4['charter'] = S4['charter'].map({'bank': 'Bank', 'cu': 'Credit Union'}); S4['main_office'] = S4['main_office'].map({True: 'Y', False: 'N'})
S4 = S4.sort_values(['charter', 'lender_name', 'main_office'], ascending=[True, True, False])

# ---- sheet 0: data dictionary
dd = [
 ('Scope', 'Banks: every FDIC-insured institution filing a Call Report (FFIEC 031/041/051), via FDIC BankFind Suite API (api.fdic.gov/banks/financials, institutions, locations, sod). Credit unions: every federally insured CU filing NCUA Form 5300, via NCUA quarterly zip files. Quarters: 2016Q1 through %s. Latest sheet = active institutions only (banks: FDIC ACTIVE=1 as of today; CUs: filed the latest quarter). Insured U.S. branches of foreign banks excluded.' % S1['Period'].iloc[0]),
 ('Units', 'Dollar columns in $ millions. Percent columns stored as fractions and formatted as %. FDIC reports thousands (converted); NCUA reports whole dollars.'),
 ('Period convention', 'Each year shows the same quarter as the latest available quarter (currently Q2 = June 30). YTD items (net change, granted, income-based ratios) are therefore 6-month figures each year, annualized where noted.'),
 ('Total Deposits', 'Bank: DEP (total deposits incl. foreign offices). CU: Acct_018 total shares and deposits.'),
 ('Total Assets', 'Bank: ASSET. CU: Acct_010.'),
 ('Total Loans', 'Bank: LNLSGR gross loans and leases. CU: Acct_025B.'),
 ('Loan to Deposit', 'Gross loans / total deposits.'),
 ('Tier 1 Leverage', 'Bank: RBC1AAJ as reported (tier 1 capital / average assets). CU: net worth / total assets (the NCUA net-worth ratio is the CU analog; CUs have no separate tier 1 measure).'),
 ('Net Worth Ratio', 'Bank: total equity capital / total assets (book). CU: Acct_997 net worth / Acct_010 assets (regulatory PCA ratio).'),
 ('Funding Cost', 'Annualized YTD total interest expense / average (deposits + borrowings), average = mean of prior Dec 31 and quarter-ends YTD. Same formula both charters. Bank borrowings = fed funds purchased + other borrowed funds + subordinated debt; CU = Acct_860C. FDIC-published column = INTEXPY (interest expense / avg earning assets) for reference.'),
 ('NIM', 'Annualized YTD net interest income / average total assets, same formula both charters. FDIC-published NIMY (on average earning assets) shown alongside for banks; it runs ~0.2-0.5 pt higher than the avg-assets basis.'),
 ('Investment CRE as % of Capital', '(Non-owner-occupied nonfarm nonresidential + multifamily + construction & land development) / capital base. Bank capital base = total risk-based capital (RBC); banks electing the Community Bank Leverage Ratio report no RBC, so tier 1 capital + loan-loss allowance is used (examiner convention). CU capital base = net worth. The 2006 interagency guidance flags banks above 300% (and above 100% for construction); no equivalent CU threshold exists.'),
 ('CRE components', 'Bank: LNRENROT (non-OO nonfarm nonres), LNRENROW (OO), LNREMULT, LNRECONS (all construction & land incl. 1-4 family construction). CU 2017Q2 onward (FS220L schedule): Acct_400J2+400J3, 400H2+400H3, 400M+400M1, 143B3+143B4 (member + purchased participations). CU 2016Q1-2017Q1 (old form): 400J, 400H, 400G (non-farm residential business loans - includes multifamily AND 1-4 family investor property; shown in the Multifamily column), 143B. Farmland excluded from CRE.'),
 ('Brokered Deposits %', 'Bank: BRO / DEP. CU: NCUA collects no brokered-deposit dollar amount - only a yes/no flag (Acct_879T, shown in "Uses brokered deposits"). Non-member deposits % (Acct_880 / Acct_018) is shown as the closest CU analog.'),
 ('CRE Net Change YTD', 'Total CRE balance now minus balance at prior Dec 31. Available for both charters; understates originations by payoffs/paydowns/sales. Investment-CRE version excludes owner-occupied.'),
 ('CRE Granted YTD (CU only)', 'Acct_475K2 (2017Q2+) / Acct_475K (earlier): dollar amount of real-estate-secured member commercial loans granted or purchased YTD; count from Acct_090K2 / 090K. Avg loan = $ / count. Banks do not report originations.'),
 ('Remaining CRE Cap (bank)', '3.0 x capital base - investment CRE. Remaining construction cap = 1.0 x capital base - construction. Negative = above the guidance level.'),
 ('MBL Cap (CU)', 'Statutory member business loan cap = lesser of 1.75 x net worth or 12.25% of assets (1.75 x 7%). Used % = Acct_400A net MBL balance / cap. Low-income-designated CUs (LIMITED_INC=1) are exempt and shown blank. Pre-2017 Acct_400A used a broader definition.'),
 ('Reported CRE Interest Rate (CU only)', 'Acct_525 "Interest rate of commercial loans/lines of credit, real estate secured" (reported in basis points on the 5300 loan schedule; collected since 2017Q3). Banks report no loan-level rates.'),
 ('Blended Loan Yield', 'Annualized YTD interest & fee income on loans / average gross loans. Same formula both charters.'),
 ('Non-Performing Loans %', 'Bank: noncurrent loans (90+ days past due + nonaccrual) / gross loans = FDIC NCLNLSR. CU: Acct_041B loans delinquent 60+ days / Acct_025B. Definitions differ by regulator; shown as reported.'),
 ('CRE Non-Performing % (bank)', '(noncurrent nonfarm nonres + multifamily + construction) / total CRE.'),
 ('ROA', 'Annualized YTD net income / average assets.'),
 ('Est. Legal Lending Limit', '15% of capital base. National banks: 15% of capital & surplus unsecured (state limits vary 15-25%). CUs: 15% of net worth to one borrower (12 CFR 723). Approximation only.'),
 ('Mergers', 'History is as reported by each charter; an acquired institution keeps its own pre-merger history under its own ID. Successor IDs are in the lender directory.'),
 ('Branches sheet', 'Illinois only, as a sample. Bank branches: FDIC locations with Summary of Deposits branch deposits for each June 30, 2016-2026. CU branches: NCUA branch file (addresses only - NCUA publishes no branch deposits; coordinates to be geocoded).'),
]
S0 = pd.DataFrame(dd, columns=['Item', 'Definition / source'])

with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    S0.to_excel(xw, sheet_name='Data dictionary', index=False)
    S1.to_excel(xw, sheet_name=f"Latest {S1['Period'].iloc[0]} all lenders", index=False)
    cov.to_excel(xw, sheet_name='Coverage by quarter', index=False)
    S4.to_excel(xw, sheet_name='Branches IL sample', index=False)

wb = load_workbook(OUT)
kinds = [k for _, _, k in COLS]
ws = wb['Data dictionary']; style(ws, ['txt', 'txt'], freeze='A2', widths={1: 34, 2: 160})
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
style(wb[f"Latest {S1['Period'].iloc[0]} all lenders"], kinds, widths={2: 34})
style(wb['Coverage by quarter'], ['txt', 'txt', 'num', 'num', 'num', 'num', 'num', 'num'], freeze='A2')
for c in ['D', 'E', 'F', 'G']:
    for cell in wb['Coverage by quarter'][c][1:]: cell.number_format = '#,##0.0'
style(wb['Branches IL sample'], ['txt'] * 12 + ['num', 'num'] + ['musd'] * (S4.shape[1] - 14), freeze='D2', widths={3: 30, 5: 30})
for c in ['M', 'N']:
    for cell in wb['Branches IL sample'][c][1:]: cell.number_format = '0.0000'
wb.save(OUT)
print(OUT, {n: wb[n].max_row for n in wb.sheetnames})
# history workbook: core columns only
core = ['Lender ID', 'Name', 'Charter', 'State', 'Period', 'Total Deposits ($M)', 'Total Assets ($M)', 'Loan to Deposit (%)', 'Tier 1 Leverage (%)', 'Net Worth Ratio (%)',
        'Funding Cost (%) [int exp / avg deposits+borrowings]', 'NIM (%) [NII / avg assets]', 'Investment CRE as % of Capital', 'Total Investment CRE ($M) [nonOO+multi+constr]', 'Total CRE ($M)',
        'Brokered Deposits % (banks)', 'Non-member Deposits % (CUs)', 'CRE Net Change YTD ($M) [total CRE]', 'CRE Granted YTD ($M) [CU only]', 'CRE Granted YTD Avg Loan ($M) [CU only]',
        'Remaining CRE Cap ($M) [bank: 300% of capital]', 'Remaining MBL Cap ($M) [CU, non-LICU]', 'Reported CRE Interest Rate (%) [CU only]', 'Blended Loan Yield (%)',
        'Non-Performing Loans % [bank 90+/nonaccrual; CU 60+]', 'ROA (%)']
ck = {n: k for n, _, k in COLS}
with pd.ExcelWriter(OUT2, engine='openpyxl') as xw:
    S0.to_excel(xw, sheet_name='Data dictionary', index=False)
    S2[core].to_excel(xw, sheet_name=f'History Q{lq} 2016-{latest.year}', index=False)
wb = load_workbook(OUT2)
ws = wb['Data dictionary']; style(ws, ['txt', 'txt'], freeze='A2', widths={1: 34, 2: 160})
style(wb[f'History Q{lq} 2016-{latest.year}'], [ck[c] for c in core], widths={2: 34})
wb.save(OUT2); print(OUT2, {n: wb[n].max_row for n in wb.sheetnames})
